import os

import pandas as pd

import docx

import openpyxl

import pyodbc  # For working with Access databases

from concurrent.futures import ProcessPoolExecutor



def search_in_csv_chunk(file_path, search_term):

    chunk_size = 10000  # Adjust as needed

    df_chunks = pd.read_csv(file_path, chunksize=chunk_size)

    results = []

    for chunk in df_chunks:

        filtered_chunk = chunk[chunk.apply(lambda row: search_term in row.astype(str).values, axis=1)]

        results.extend(filtered_chunk.values.tolist())

    return results



def search_in_txt_chunk(file_path, search_term):

    results = []

    with open(file_path, 'r') as file:

        for line in file:

            if search_term in line:

                results.append(line.strip())

    return results



def search_in_docx_chunk(file_path, search_term):

    doc = docx.Document(file_path)

    results = []

    for para in doc.paragraphs:

        if search_term in para.text:

            results.append(para.text)

    return results



def search_in_xlsx_chunk(file_path, search_term):

    wb = openpyxl.load_workbook(file_path, read_only=True)

    results = []

    for sheet in wb:

        for row in sheet.iter_rows(values_only=True):

            if any(search_term in str(cell) for cell in row):

                results.append(row)

    return results



def search_in_accdb_chunk(file_path, search_term):

    conn_str = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};Dbq={file_path};"

    with pyodbc.connect(conn_str) as conn:

        cursor = conn.cursor()

        query = f"SELECT * FROM YourTableName WHERE YourColumnName LIKE '%{search_term}%'"

        cursor.execute(query)

        results = cursor.fetchall()

    return results



def search_files_parallel(directory, search_term):

    results = {}

    with ProcessPoolExecutor() as executor:

        for root, _, files in os.walk(directory):

            for file in files:

                file_path = os.path.join(root, file)

                if file.endswith('.csv'):

                    results[file] = executor.submit(search_in_csv_chunk, file_path, search_term)

                elif file.endswith('.txt'):

                    results[file] = executor.submit(search_in_txt_chunk, file_path, search_term)

                elif file.endswith('.docx'):

                    results[file] = executor.submit(search_in_docx_chunk, file_path, search_term)

                elif file.endswith('.xlsx'):

                    results[file] = executor.submit(search_in_xlsx_chunk, file_path, search_term)

                elif file.endswith('.accdb'):

                    results[file] = executor.submit(search_in_accdb_chunk, file_path, search_term)

    return results



# Example usage

directory = 'directory'

search_term = 'search_term'

search_results = search_files_parallel(directory, search_term)



for file, future in search_results.items():

    result = future.result()

    print(f"Results in {file}:")

    for item in result:

        print(item)
